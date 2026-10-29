"""
Escritura de CPE en el Excel de avance de cosecha.

- Trabaja en memoria: recibe el Excel como bytes (o ruta) y devuelve el
  Excel actualizado como bytes, para que la app web lo ofrezca como descarga.
- Detecta las hojas de CPE (las que tienen columnas CTG y CPE en la fila 1).
- Ubica cada columna por el nombre del encabezado, no por la posición
  (en tus hojas el orden cambia: en 'CPE Soja1' va CPE antes que CTG).
- Sugiere la hoja destino según el grano y los CTG más cercanos.
- Avisa si el CTG ya está cargado.
- Agrega las filas al final copiando el formato de la fila de arriba.
"""
from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path

import openpyxl

# Nombre lógico -> textos posibles del encabezado (normalizados)
SINONIMOS = {
    "cultivo": ["cultivo"],
    "campo": ["campo"],
    "lote": ["lote"],
    "kgs": ["kgs neto", "kg neto", "kilos neto", "kilos", "kgs", "neto"],
    "ctg": ["nro ctg", "n ctg", "ctg"],
    "cpe": ["cpe", "nro cpe", "n cpe"],
    "patente": ["patente", "dominio", "patentes"],
    "fecha": ["fecha"],
    "transportista": ["transportista", "transporte", "empresa transportista"],
}
COLUMNAS_EXTRA = ["patente", "fecha", "transportista"]
ETIQUETAS = {
    "cultivo": "Cultivo", "campo": "Campo", "lote": "Lote", "kgs": "Kgs Neto",
    "ctg": "NRO CTG", "cpe": "CPE", "patente": "Patente", "fecha": "Fecha",
    "transportista": "Transportista",
}
GRANOS = {"maiz": "maiz", "soja": "soja", "girasol": "girasol", "trigo": "trigo",
          "sorgo": "sorgo", "cebada": "cebada"}

Fuente = str | Path | bytes


def normalizar(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-z0-9 ]+", " ", t.lower())
    return re.sub(r"\s+", " ", t).strip()


def _ctg_num(v) -> int | None:
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def _abrir(fuente: Fuente) -> openpyxl.Workbook:
    """Abre el libro desde bytes (subido por la web) o desde una ruta (uso local / tests)."""
    if isinstance(fuente, (bytes, bytearray)):
        fuente = io.BytesIO(fuente)
    return openpyxl.load_workbook(fuente)  # con fórmulas, sin valores calculados


@dataclass
class HojaCPE:
    nombre: str
    columnas: dict[str, int]                 # nombre lógico -> nro de columna (1-based)
    ultima_fila: int                         # última fila con datos
    cultivo_etiqueta: str = ""               # valor más usado en la columna Cultivo
    campos: list[str] = field(default_factory=list)
    lotes_por_campo: dict[str, list[str]] = field(default_factory=dict)
    ctgs: list[int] = field(default_factory=list)
    ctg_como_texto: bool = False

    @property
    def faltantes(self) -> list[str]:
        return [c for c in COLUMNAS_EXTRA if c not in self.columnas]


@dataclass
class Libro:
    hojas: dict[str, HojaCPE]
    indice_ctg: dict[int, list[tuple[str, int]]]   # ctg -> [(hoja, fila)]

    def duplicados(self, ctg) -> list[tuple[str, int]]:
        n = _ctg_num(ctg)
        return self.indice_ctg.get(n, []) if n else []

    def sugerir_hoja(self, grano: str, ctg) -> str | None:
        clave = next((v for k, v in GRANOS.items() if k in normalizar(grano)), normalizar(grano))
        candidatas = [h for h in self.hojas.values()
                      if clave and (clave in normalizar(h.nombre) or clave in normalizar(h.cultivo_etiqueta))]
        if not candidatas:
            candidatas = list(self.hojas.values())
        if not candidatas:
            return None
        n = _ctg_num(ctg) or 0
        # La hoja "activa" es la que tiene más CTG cercanos (los CTG son correlativos en el tiempo)
        def puntaje(h: HojaCPE):
            cercanos = sum(1 for x in h.ctgs if abs(x - n) <= 2_000_000)
            return (cercanos, max(h.ctgs, default=0))
        return max(candidatas, key=puntaje).nombre


def _mapear_encabezados(ws) -> dict[str, int]:
    cols: dict[str, int] = {}
    for celda in ws[1]:
        h = normalizar(celda.value)
        if not h:
            continue
        for logico, opciones in SINONIMOS.items():
            if logico not in cols and h in opciones:
                cols[logico] = celda.column
                break
    return cols


def _ultima_fila(ws, cols: dict[str, int]) -> int:
    claves = [cols[k] for k in ("cultivo", "campo", "kgs", "ctg", "cpe") if k in cols]
    for fila in range(ws.max_row, 1, -1):
        if any(ws.cell(fila, c).value not in (None, "") for c in claves):
            return fila
    return 1


def analizar_libro(fuente: Fuente) -> Libro:
    wb = _abrir(fuente)
    hojas: dict[str, HojaCPE] = {}
    indice: dict[int, list[tuple[str, int]]] = defaultdict(list)

    for ws in wb.worksheets:
        cols = _mapear_encabezados(ws)
        if "ctg" not in cols or "cpe" not in cols:
            continue
        ultima = _ultima_fila(ws, cols)
        cultivos, campos = Counter(), Counter()
        lotes: dict[str, Counter] = defaultdict(Counter)
        ctgs, textos = [], 0
        for f in range(2, ultima + 1):
            v_ctg = ws.cell(f, cols["ctg"]).value
            n = _ctg_num(v_ctg)
            if n:
                ctgs.append(n)
                indice[n].append((ws.title, f))
                textos += isinstance(v_ctg, str)
            if "cultivo" in cols and ws.cell(f, cols["cultivo"]).value:
                cultivos[str(ws.cell(f, cols["cultivo"]).value).strip()] += 1
            campo = ws.cell(f, cols["campo"]).value if "campo" in cols else None
            if campo:
                campo = str(campo).strip()
                campos[campo] += 1
                if "lote" in cols:
                    lote = ws.cell(f, cols["lote"]).value
                    if isinstance(lote, (int, float)) and not isinstance(lote, bool):
                        lote = int(lote) if float(lote).is_integer() else lote
                    if lote not in (None, "") and not isinstance(lote, datetime):
                        lotes[campo][str(lote)] += 1
        hojas[ws.title] = HojaCPE(
            nombre=ws.title, columnas=cols, ultima_fila=ultima,
            cultivo_etiqueta=cultivos.most_common(1)[0][0] if cultivos else "",
            campos=[c for c, _ in campos.most_common()],
            lotes_por_campo={c: [l for l, _ in cnt.most_common()] for c, cnt in lotes.items()},
            ctgs=ctgs, ctg_como_texto=textos > len(ctgs) / 2,
        )
    wb.close()
    return Libro(hojas=hojas, indice_ctg=dict(indice))


def campo_canonico(texto: str, conocidos: list[str]) -> str:
    """Devuelve el campo tal como está escrito en el Excel si coincide (sin importar mayúsculas)."""
    n = normalizar(texto)
    for c in conocidos:
        if normalizar(c) == n:
            return c
    for c in conocidos:  # coincidencia parcial: 'Paunero' -> 'Fortin Paunero'
        if n and (n in normalizar(c) or normalizar(c) in n):
            return c
    return texto.strip()


def transportista_corto(nombre: str, cuit: str = "") -> str:
    """En tus planillas usan el apellido ('Barbesini'). Persona física (CUIT 20/23/24/27) -> apellido."""
    if cuit[:2] in {"20", "23", "24", "27"} and nombre:
        return nombre.split()[0].title()
    return nombre.title()


def _valor_lote(lote: str):
    lote = str(lote).strip()
    return int(lote) if lote.isdigit() else lote


@dataclass
class FilaNueva:
    hoja: str
    cultivo: str
    campo: str
    lote: str
    kgs: int
    ctg: str
    cpe: str
    patente: str = ""
    fecha: date | None = None
    transportista: str = ""


def guardar_filas(fuente: Fuente, filas: list[FilaNueva], agregar_columnas_faltantes: bool = False) -> dict:
    """Agrega las filas al final de cada hoja. No toca la fuente: devuelve el Excel nuevo en 'contenido'."""
    wb = _abrir(fuente)
    escritas: list[tuple[str, int]] = []
    avisos: list[str] = []

    for fila in filas:
        ws = wb[fila.hoja]
        cols = _mapear_encabezados(ws)
        if agregar_columnas_faltantes:
            for extra in COLUMNAS_EXTRA:
                if extra not in cols:
                    nueva = max([c.column for c in ws[1] if c.value not in (None, "")] + [0]) + 1
                    ws.cell(1, nueva, ETIQUETAS[extra])
                    ws.cell(1, nueva)._style = copy(ws.cell(1, max(cols.values()))._style)
                    cols[extra] = nueva
        destino = _ultima_fila(ws, cols) + 1
        modelo = destino - 1 if destino > 2 else None

        ctgs_texto = sum(isinstance(ws.cell(f, cols["ctg"]).value, str) for f in range(2, destino))
        ctg_valor = fila.ctg if ctgs_texto > (destino - 2) / 2 else _ctg_num(fila.ctg)

        valores = {
            "cultivo": fila.cultivo, "campo": fila.campo, "lote": _valor_lote(fila.lote),
            "kgs": fila.kgs, "ctg": ctg_valor, "cpe": fila.cpe,
            "patente": fila.patente or None,
            "fecha": datetime.combine(fila.fecha, datetime.min.time()) if fila.fecha else None,
            "transportista": fila.transportista or None,
        }
        for clave, valor in valores.items():
            if clave not in cols:
                if valor not in (None, "") and clave in COLUMNAS_EXTRA:
                    avisos.append(f"'{fila.hoja}' no tiene columna {ETIQUETAS[clave]}: no se cargó.")
                continue
            celda = ws.cell(destino, cols[clave])
            if modelo:  # copiar formato de la fila de arriba
                celda._style = copy(ws.cell(modelo, cols[clave])._style)
            celda.value = valor
            if clave == "fecha" and valor is not None:
                celda.number_format = "dd/mm/yyyy"
            if clave == "ctg" and isinstance(valor, int):
                celda.number_format = "0"
        escritas.append((fila.hoja, destino))

    salida = io.BytesIO()
    try:
        wb.save(salida)
    finally:
        wb.close()
    return {"contenido": salida.getvalue(), "filas": escritas, "avisos": sorted(set(avisos))}
