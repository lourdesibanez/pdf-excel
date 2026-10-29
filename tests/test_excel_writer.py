from pathlib import Path
from datetime import date
import io
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from excel_writer import FilaNueva, analizar_libro, guardar_filas


def libro_ejemplo() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CPE Maiz Tard."
    ws.append(["Cultivo", "Campo", "Lote", "Kgs Neto", "NRO CTG", "CPE"])
    ws.append(["Maiz Tardio", "Fortin Paunero", 3, 30000, 10134300001, "00001-00026500"])
    soja = wb.create_sheet("CPE Soja1")
    soja.append(["Cultivo", "Campo", "Lote", "Kgs Neto", "CPE", "CTG", "Patente", "Fecha", "Transportista"])
    soja.append(["Soja", "La Esperanza", "5", 28000, "00001-00020000", "10100000000", "AAA111", None, "Perez"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fila(lote: str, kgs: int) -> FilaNueva:
    return FilaNueva(hoja="CPE Maiz Tard.", cultivo="Maiz Tardio", campo="Fortin Paunero", lote=lote, kgs=kgs,
                     ctg="10134300029", cpe="00001-00026528", patente="HDE283", fecha=date(2026, 8, 11),
                     transportista="Barbesini Claudio Oscar")


def test_analizar_desde_bytes():
    libro = analizar_libro(libro_ejemplo())
    assert set(libro.hojas) == {"CPE Maiz Tard.", "CPE Soja1"}
    assert libro.sugerir_hoja("Maíz", "10134300029") == "CPE Maiz Tard."
    assert libro.duplicados("10134300001") == [("CPE Maiz Tard.", 2)]
    assert libro.hojas["CPE Maiz Tard."].faltantes == ["patente", "fecha", "transportista"]
    assert libro.hojas["CPE Soja1"].ctg_como_texto


def test_guardar_devuelve_excel_nuevo():
    original = libro_ejemplo()
    res = guardar_filas(original, [fila("3", 20000), fila("4", 9960)])
    assert res["filas"] == [("CPE Maiz Tard.", 3), ("CPE Maiz Tard.", 4)]
    assert len(res["avisos"]) == 3  # la hoja no tiene Patente/Fecha/Transportista

    nuevo = analizar_libro(res["contenido"])
    assert nuevo.duplicados("10134300029") == [("CPE Maiz Tard.", 3), ("CPE Maiz Tard.", 4)]
    assert analizar_libro(original).duplicados("10134300029") == []  # el original no cambia

    ws = openpyxl.load_workbook(io.BytesIO(res["contenido"]))["CPE Maiz Tard."]
    assert [c.value for c in ws[4]] == ["Maiz Tardio", "Fortin Paunero", 4, 9960, 10134300029, "00001-00026528"]


def test_agregar_columnas_faltantes():
    res = guardar_filas(libro_ejemplo(), [fila("3", 29960)], agregar_columnas_faltantes=True)
    assert res["avisos"] == []
    nuevo = analizar_libro(res["contenido"])
    assert nuevo.hojas["CPE Maiz Tard."].faltantes == []
